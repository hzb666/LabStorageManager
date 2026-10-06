from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook
from starlette.requests import Request

from app.api import inventory, inventory_extended_routes
from app.services import export_batch
from app.services.export_batch import BatchFetchResult
from app.services.xlsx_export import export_inventory_xlsx


@pytest.mark.parametrize(
    ("location", "keeper_id", "keeper_names", "expected"),
    [
        (None, 7, {7: "张三"}, "张三暂存"),
        ("", 7, {7: "张三"}, "张三暂存"),
        ("A柜", 7, {7: "张三"}, "A柜"),
        (None, None, {}, None),
        (None, 7, {}, "未知用户暂存"),
    ],
)
@pytest.mark.asyncio
async def test_inventory_workbook_location(location, keeper_id, keeper_names, expected):
    response = export_inventory_xlsx(
        [{"storage_location": location, "temporary_keeper_id": keeper_id}],
        temporary_keeper_names=keeper_names,
    )
    content = b"".join([chunk async for chunk in response.body_iterator])
    workbook = load_workbook(BytesIO(content))
    try:
        sheet = workbook["库存"]
        assert sheet.cell(1, 8).value == "位置"
        assert sheet.cell(2, 8).value == expected
    finally:
        workbook.close()


def test_inventory_export_resolves_keeper_names_in_batch(monkeypatch):
    endpoint = next(
        route.endpoint
        for route in inventory.router.routes
        if route.path.endswith("/export")
    )
    items = [SimpleNamespace(temporary_keeper_id=7), SimpleNamespace(temporary_keeper_id=None)]
    result = BatchFetchResult(items=items, total_count=2, is_truncated=False)
    monkeypatch.setattr(export_batch, "batch_fetch_all", lambda *_: result)
    monkeypatch.setattr(inventory_extended_routes, "enforce_export_rate_limit", lambda *_: None)
    monkeypatch.setattr(inventory_extended_routes, "log_inventory_export_operation", lambda *_, **__: None)
    resolve_names = MagicMock(return_value={7: "张三"})
    export_workbook = MagicMock()
    monkeypatch.setattr(inventory_extended_routes, "batch_get_user_names", resolve_names)
    monkeypatch.setattr(inventory_extended_routes, "export_inventory_xlsx", export_workbook)
    db = MagicMock()
    request = Request({"type": "http", "method": "GET", "path": "/export", "headers": []})

    endpoint(request, SimpleNamespace(id=1), db)

    resolve_names.assert_called_once_with(db, {7})
    export_workbook.assert_called_once_with(items, temporary_keeper_names={7: "张三"})
